"""
Fills ENTRY_EPPO, ENTRY_MOPS and ENTRY_OILFUND_SUSTAINABILITY automatically, replacing the
manual daily data-entry step. Runs before main.py in the GitHub Actions workflow.

Sources (see project chat/README for how each was reverse-engineered from real data):
  - EPPO price structure: a predictable daily .xlsx at eppo.go.th (no scraping needed).
  - MOPS: reused from the OIL BOT project's own MopsLog tab (already captured there).
  - Oil Fund Sustainability: the official weekly PDF at offo.or.th when a new one is
    published; otherwise rolled forward from the previous day using that day's EPPO-derived
    Net Fund Impact (Total_Balance) / Daily_Subsidy (Cash_Remaining).

Auth: a single Google service account (SA_EMAIL / SA_PRIVATE_KEY env vars) needs Editor
access to the procurement_decision sheet AND read access to the OIL BOT Database sheet.
"""
import io
import os
import re
from datetime import datetime, timedelta, date
from zoneinfo import ZoneInfo

import openpyxl
import pdfplumber
import requests

from common import (
    SHEET_ID, OILBOT_SHEET_ID, TABS, get_token, sheets_get, sheets_append, sheets_update, sheets_replace_tab,
    col_to_dicts, excel_serial_to_date, fmt_date_us,
)

TZ = ZoneInfo("Asia/Bangkok")


def _key_str(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v).strip()


# ============================================================
# EPPO PRICE STRUCTURE
# ============================================================
def eppo_xlsx_url(d: date) -> str:
    return f"https://www.eppo.go.th/wp-content/uploads/{d.year}/{d.month:02d}/pt-price-st-{d.year}-{d.month}-{d.day}.xlsx"


def fetch_eppo_workbook(d: date):
    url = eppo_xlsx_url(d)
    r = requests.get(url, timeout=30)
    if r.status_code != 200:
        return None
    return openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)


def parse_eppo_rows(wb, d: date):
    # ชื่อแท็บเปลี่ยนรูปแบบไปมา (เจอทั้ง "Oil Price Structure" เฉยๆ และ "Oil Price Structure-รวมภาษี")
    # จับคู่แบบ "ขึ้นต้นด้วย" แทนเทียบตรงตัว
    sheet_name = next((n for n in wb.sheetnames if n.strip().startswith("Oil Price Structure")), None)
    if sheet_name is None:
        raise RuntimeError(f"EPPO xlsx for {d}: no 'Oil Price Structure*' tab found, found {wb.sheetnames}")
    ws = wb[sheet_name]
    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and str(row[1] or "").strip().upper().startswith("UNIT"):
            header_row_idx = i
            break
    if header_row_idx is None:
        raise RuntimeError("EPPO xlsx: header row not found")

    out = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        name = str(row[1] or "").strip()
        if not name:
            break  # blank row = end of the fuel table
        ex_refin, _discount, excise, mtax, oilfund, consv, ws_price, vat_ws, ws_vat, mm, vat_mm, retail = row[2:14]
        out.append({
            "oil_type": name,
            "ex_refin": ex_refin, "excise": excise, "mtax": mtax, "oilfund": oilfund,
            "consv": consv, "wholesale": ws_price, "vat_ws": vat_ws, "ws_vat": ws_vat,
            "mm": mm, "vat_mm": vat_mm, "retail": retail,
        })
    return out


def load_consumption_ref(tok):
    rows = sheets_get(tok, SHEET_ID, f"{TABS['NATIONAL_CONSUMPTION_REF']}!A1:H20")
    ref = {}
    for r in col_to_dicts(rows):
        key = str(r.get("EPPO Oil Type Key") or "").strip()
        avg = r.get("Avg by Key (ล้านลิตร/วัน)")
        if key and isinstance(avg, (int, float)):
            ref[key] = float(avg)
    return ref


# คอลัมน์ B..M ของ Entry_Eppo ตามลำดับจริงในชีท (ไม่รวม Date/Avg Consumption/Estimated */Obs Rank
# ซึ่งคำนวณ/จัดการแยกต่างหากตอนสร้างแถว)
EPPO_PRICE_COLS = [
    "oil type", "EX-REFIN.", "EXCISE TAX", "M. TAX", "OIL FUND", "CONSV. FUND",
    "WHOLESALE (WS)", "VAT (WS)", "WS&VAT", "MARKETING MARGIN", "VAT (MM)", "RETAIL",
]


def ensure_eppo_source_header(tok):
    """คอลัมน์ S 'Source' บอกว่าแถวนี้มาจากไฟล์ EPPO จริง (FRESH) หรือลากค่าเดิมมา (CARRIED) —
    main.py ใช้ตรวจว่าค่า ExRef/กองทุนใน dashboard เก่าแค่ไหน (เดิมแยกไม่ออก จึงตัดสินใจจากค่าค้างโดยไม่รู้ตัว)"""
    head = sheets_get(tok, SHEET_ID, f"{TABS['ENTRY_EPPO']}!S1:S1")
    if not head or not head[0] or head[0][0] != "Source":
        sheets_update(tok, SHEET_ID, f"{TABS['ENTRY_EPPO']}!S1", [["Source"]])


def carry_forward_eppo(tok, existing_dicts, obs_count, date_str):
    """ไม่มีไฟล์ EPPO วันนี้ (เช่น วันหยุด/ยังไม่เผยแพร่) — ตามหลักที่ตกลงไว้ ให้ใช้ตัวเลขของ
    วันก่อนหน้าซ้ำแทนที่จะปล่อยช่องว่าง (ราคายังไม่เปลี่ยนจนกว่าจะมีประกาศใหม่จริง)"""
    dated = [(excel_serial_to_date(r.get("Date")), r) for r in existing_dicts if r.get("Date") is not None]
    if not dated:
        print("[EPPO] no previous rows to carry forward from, skipping")
        return None
    latest_date = max(d for d, _ in dated)
    latest_rows = [r for d, r in dated if d == latest_date]

    new_rows = []
    day_net_impact = 0.0
    for r in latest_rows:
        ot = r.get("oil type")
        obs_count[ot] = obs_count.get(ot, 0) + 1
        subsidy = r.get("Estimated Subsidy (mn baht/day)") or 0
        collection = r.get("Estimated Collection (mn baht/day)") or 0
        net_impact = round(collection - subsidy, 2)
        day_net_impact += net_impact
        row = [date_str] + [r.get(c) for c in EPPO_PRICE_COLS]
        row += [r.get("Avg Consumption (mn L/day)"), subsidy, collection, net_impact, obs_count[ot], "CARRIED"]
        new_rows.append(row)

    ensure_eppo_source_header(tok)
    sheets_append(tok, SHEET_ID, f"{TABS['ENTRY_EPPO']}!A1:S1", new_rows)
    print(f"[EPPO] carried forward {len(new_rows)} rows from {latest_date} as {date_str}")
    return {"date_str": date_str, "net_impact": day_net_impact,
            "subsidy": sum(r.get("Estimated Subsidy (mn baht/day)") or 0 for r in latest_rows),
            "collection": sum(r.get("Estimated Collection (mn baht/day)") or 0 for r in latest_rows)}


def ingest_eppo(tok, today: date):
    existing = sheets_get(tok, SHEET_ID, f"{TABS['ENTRY_EPPO']}!A1:Q5000")
    existing_dicts = col_to_dicts(existing)
    # Date เก็บเป็น serial number ของ Sheets ไม่ใช่ string — ต้องแปลงก่อนเทียบ ไม่งั้นจะไม่มีวันตรงกัน
    seen_dates = {excel_serial_to_date(r.get("Date")) for r in existing_dicts if r.get("Date") is not None}
    obs_count = {}
    for r in existing_dicts:
        ot = r.get("oil type")
        if ot:
            obs_count[ot] = obs_count.get(ot, 0) + 1

    date_str = fmt_date_us(today)
    if today in seen_dates:
        print(f"[EPPO] {date_str} already recorded, skipping")
        return None

    wb = fetch_eppo_workbook(today)
    if wb is None:
        print(f"[EPPO] {eppo_xlsx_url(today)} not published yet — carrying forward yesterday's values")
        return carry_forward_eppo(tok, existing_dicts, obs_count, date_str)

    consumption = load_consumption_ref(tok)
    fuel_rows = parse_eppo_rows(wb, today)

    new_rows = []
    day_net_impact = 0.0
    for f in fuel_rows:
        avg = consumption.get(f["oil_type"])
        oilfund = f["oilfund"] if isinstance(f["oilfund"], (int, float)) else None
        if avg and oilfund is not None:
            net = oilfund * avg
            subsidy = max(0.0, -net)
            collection = max(0.0, net)
            day_net_impact += net
        else:
            subsidy = collection = 0.0
            avg = avg or ""
        obs_count[f["oil_type"]] = obs_count.get(f["oil_type"], 0) + 1
        new_rows.append([
            date_str, f["oil_type"], f["ex_refin"], f["excise"], f["mtax"], f["oilfund"],
            f["consv"], f["wholesale"], f["vat_ws"], f["ws_vat"], f["mm"], f["vat_mm"], f["retail"],
            avg, round(subsidy, 2), round(collection, 2), round(collection - subsidy, 2), obs_count[f["oil_type"]], "FRESH",
        ])

    ensure_eppo_source_header(tok)
    sheets_append(tok, SHEET_ID, f"{TABS['ENTRY_EPPO']}!A1:S1", new_rows)
    print(f"[EPPO] appended {len(new_rows)} rows for {date_str}, day net impact = {day_net_impact:.2f}")
    return {"date_str": date_str, "net_impact": day_net_impact,
            "subsidy": sum(r[14] for r in new_rows), "collection": sum(r[15] for r in new_rows)}


# ============================================================
# MOPS (reused from OIL BOT's own MopsLog)
# ============================================================
def ingest_mops(tok, today: date):
    mops_rows = col_to_dicts(sheets_get(tok, OILBOT_SHEET_ID, "MopsLog!A1:J20000"))
    by_date = {}
    for r in mops_rows:
        d = excel_serial_to_date(r.get("mops_date"))
        if d:
            by_date[d] = r  # later rows overwrite earlier ones for the same date (last wins)

    existing = col_to_dicts(sheets_get(tok, SHEET_ID, f"{TABS['ENTRY_MOPS']}!A1:F20000"))
    # Date เก็บเป็น serial number ของ Sheets ไม่ใช่ string — ต้องแปลงก่อนเทียบ ไม่งั้นจะไม่มีวันตรงกัน
    seen = {excel_serial_to_date(r.get("Date")) for r in existing if r.get("Date") is not None}
    obs_count = {"95": 0, "DS": 0}
    # ราคาล่าสุดที่เคยบันทึกไว้ต่อชนิด — ใช้เป็นฐานถ้าวันนี้ไม่มีข้อมูลใหม่เลย (เริ่มจากของเดิมในชีท)
    last_price = {}
    row_of = {}  # (date, key) -> (เลขแถวในชีท, ราคา)
    for i, r in enumerate(existing):
        # Normalized Key ของ G95 อ่านกลับมาเป็นเลข 95 (หรือข้อความ "95" ในแถวเก่า) ต้องเทียบเป็นสตริงเสมอ
        # ไม่งั้น `95 in {"95": ..}` เป็น False แล้วไปหยิบราคาจากแถวข้อความเก่าๆ มาเป็น "ราคาล่าสุด"
        k = _key_str(r.get("Normalized Key"))
        if k in obs_count:
            obs_count[k] += 1
            last_price[k] = r.get("Price (USD/BBL)")
            d = excel_serial_to_date(r.get("Date"))
            if d:
                row_of[(d, k)] = (i + 2, r.get("Price (USD/BBL)"))

    # แถว carry-forward ของวันหนึ่งอาจถูกเขียนไว้ก่อนที่ราคาปิดจริงจะโพสต์ในไลน์ (ปิดตลาดโพสต์เย็น แต่ cron รันตั้งแต่เช้า)
    # ให้ MopsLog เป็นแหล่งความจริง: วันที่มีใน MopsLog ใช้ค่าจริง วันที่ไม่มีใช้ค่าของวันก่อนหน้า (ตามหลัก carry-forward)
    # ทำเฉพาะวันตั้งแต่ที่ MopsLog เริ่มมีข้อมูลเป็นต้นไป — ข้อมูลเก่าที่กรอกมือก่อนหน้านั้นไม่แตะ
    first_log = min(by_date) if by_date else None
    field_of = {"95": "g95_price", "DS": "diesel_price"}
    expected = {}
    for (d, k), (row_no, price) in sorted(row_of.items()):
        if first_log is None or d < first_log:
            expected[k] = price
            continue
        official = by_date[d].get(field_of[k]) if d in by_date else None
        exp = official if isinstance(official, (int, float)) else expected.get(k, price)
        expected[k] = exp
        if isinstance(price, (int, float)) and abs(exp - price) > 0.005:
            sheets_update(tok, SHEET_ID, f"{TABS['ENTRY_MOPS']}!C{row_no}", [[exp]])
            print(f"[MOPS] corrected {d} {k}: {price} -> {exp}")
    for k, v in expected.items():
        last_price[k] = v

    new_rows = []
    got_today = False
    for d in sorted(by_date):
        if d in seen:
            continue
        date_str = fmt_date_us(d)
        r = by_date[d]
        obs_count["95"] += 1
        obs_count["DS"] += 1
        last_price["95"], last_price["DS"] = r.get("g95_price"), r.get("diesel_price")
        new_rows.append([date_str, "G95", last_price["95"], "", "95", obs_count["95"]])
        new_rows.append([date_str, "DS", last_price["DS"], "", "DS", obs_count["DS"]])
        if d == today:
            got_today = True

    if today not in seen and not got_today and last_price:
        # OIL BOT ไม่มี MOPS ของวันนี้เลย (เช่น ไม่มีคนโพสต์ในกลุ่ม LINE วันนั้น) — ใช้ราคาล่าสุด
        # ที่มีซ้ำแทนตามหลักที่ตกลงไว้ แทนที่จะปล่อยช่องว่าง
        date_str = fmt_date_us(today)
        for fuel_label, key in (("G95", "95"), ("DS", "DS")):
            if last_price.get(key) is not None:
                obs_count[key] += 1
                new_rows.append([date_str, fuel_label, last_price[key], "", key, obs_count[key]])
        print(f"[MOPS] no data for {today} — carried forward last known prices")

    sheets_append(tok, SHEET_ID, f"{TABS['ENTRY_MOPS']}!A1:F1", new_rows)
    print(f"[MOPS] appended {len(new_rows)} rows ({len(new_rows) // 2} dates)")


# ============================================================
# ฟีดสัญญาณตลาดจาก OIL BOT (Trend Mops, ประกาศ "มีปรับขึ้น/ลง", ประกาศกองทุน, MOPS ปิด, ราคา Jobber)
# ============================================================
FEED_DAYS_PRICES = 21          # เก็บราคา Jobber ย้อนหลังกี่วัน (ใช้คำนวณผลลัพธ์ของคำแนะนำใน DecisionLog)


def _feed_rows(tok, src_tab, columns):
    """อ่านแท็บจาก OIL BOT แล้วคืนแถว [หัว] + ข้อมูล เฉพาะคอลัมน์ที่ต้องการ (ตามชื่อคอลัมน์ต้นทาง)"""
    values = sheets_get(tok, OILBOT_SHEET_ID, f"{src_tab}!A1:Z100000")
    if not values:
        return [list(columns)]
    head = [str(h).strip() for h in values[0]]
    idx = [head.index(c) if c in head else None for c in columns]
    out = [list(columns)]
    for r in values[1:]:
        out.append(["" if i is None or i >= len(r) else r[i] for i in idx])
    return out


def _prices_feed(tok):
    """PriceLog -> Entry_Prices (provider, warehouse, fuel, price, effective) เฉพาะ FEED_DAYS_PRICES วันล่าสุด
    effective = วัน/เวลาที่ราคามีผล: ใช้ effective_time ถ้าระบุ (เช่น ปรับขึ้น 23:59 / เที่ยงคืน) โดยถ้าเวลาที่มีผลน้อยกว่าเวลาโพสต์เกิน 2 ชม.
    ถือว่าเป็นของวันถัดไป ไม่งั้นใช้เวลาที่โพสต์"""
    rows = col_to_dicts(sheets_get(tok, OILBOT_SHEET_ID, "PriceLog!A1:Q200000"))
    cutoff = (datetime.now(TZ) - timedelta(days=FEED_DAYS_PRICES)).strftime("%Y-%m-%d")
    out = [["provider", "warehouse", "fuel", "price", "effective"]]
    for r in rows:
        d, t = str(r.get("update_date") or ""), str(r.get("update_time") or "")[:5]
        price = r.get("price")
        if not re.match(r"^\d{4}-\d{2}-\d{2}$", d) or d < cutoff or not isinstance(price, (int, float)):
            continue
        posted = datetime.strptime(f"{d} {t or '00:00'}", "%Y-%m-%d %H:%M")
        eff_t = str(r.get("effective_time") or "")[:5]
        eff = posted
        if re.match(r"^\d{1,2}:\d{2}$", eff_t):
            cand = datetime.strptime(f"{d} {eff_t.zfill(5)}", "%Y-%m-%d %H:%M")
            eff = cand + timedelta(days=1) if cand < posted - timedelta(hours=2) else cand
        out.append([r.get("provider"), r.get("warehouse_name"), r.get("fuel_type"), price, eff.strftime("%Y-%m-%d %H:%M")])
    return out


def ingest_feeds(tok):
    """คัดลอกสัญญาณจากชีท OIL BOT มาไว้ในชีทนี้ (main.py อ่านชีทนี้ผ่าน CSV สาธารณะ ไม่มี credential)"""
    plan = [
        (TABS["FEED_INTRADAY"], lambda: _feed_rows(tok, "IntradayMops", ["msg_date", "msg_time", "kind", "b", "d", "wti", "wti_chg", "brent", "brent_chg"])),
        (TABS["FEED_HEADSUP"], lambda: _feed_rows(tok, "HeadsUp", ["msg_date", "msg_time", "provider", "direction", "scope", "stated_time", "has_rate", "rates"])),
        (TABS["FEED_FUND_NOTICE"], lambda: _feed_rows(tok, "FundNotice", ["msg_date", "msg_time", "issue_no", "effective_date", "diesel_old", "diesel_new", "benzene_old", "benzene_new"])),
        (TABS["FEED_MOPS_CLOSE"], lambda: _mops_close_feed(tok)),
        (TABS["FEED_PRICES"], lambda: _prices_feed(tok)),
    ]
    failed = []
    for title, build in plan:
        try:
            rows = build()
            sheets_replace_tab(tok, SHEET_ID, title, rows)
            print(f"[FEED] {title}: {len(rows) - 1} rows")
        except Exception as e:           # แท็บต้นทางอาจยังไม่ถูกสร้าง (worker สร้างเองตอนมีข้อความแรก) — ข้ามแล้วไปต่อ
            print(f"[FEED] {title} skipped: {e}")
            failed.append(title)
    return failed


def _mops_close_feed(tok):
    """MopsLog (OIL BOT) -> สคีมา mops_close: msg_date, msg_time, mops_date, g95, g95_chg, diesel, diesel_chg"""
    rows = col_to_dicts(sheets_get(tok, OILBOT_SHEET_ID, "MopsLog!A1:L100000"))
    out = [["msg_date", "msg_time", "mops_date", "g95", "g95_chg", "diesel", "diesel_chg"]]
    for r in rows:
        ts = str(r.get("timestamp_logged") or "")
        if len(ts) < 16:
            continue
        md = r.get("mops_date")
        md = excel_serial_to_date(md).isoformat() if isinstance(md, (int, float)) else str(md or "")[:10]
        out.append([ts[:10], ts[11:16], md, r.get("g95_price"), r.get("g95_change"), r.get("diesel_price"), r.get("diesel_change")])
    return out


# ============================================================
# OIL FUND SUSTAINABILITY
# ============================================================
def latest_offo_report():
    r = requests.get("https://www.offo.or.th/th/estimate/fuelfund-status", timeout=30)
    r.raise_for_status()
    links = re.findall(r'href="(https://www\.offo\.or\.th/sites/default/files/estimate/files/[^"]+\.(?:pdf|png|jpg))"', r.text)
    dated = []
    for url in links:
        m = re.search(r"(\d{2})-(\d{2})-(\d{2})", url)  # ...69-06-07... = yy-mm-dd (BE year)
        if m:
            yy, mm, dd = (int(x) for x in m.groups())
            try:
                dated.append((date(2500 + yy - 543, mm, dd), url))
            except ValueError:
                continue
    if not dated:
        return None, None
    dated.sort(key=lambda x: x[0])
    return dated[-1]  # (end_date, url)


def _num(cell):
    """แปลงข้อความในเซลล์ตารางเป็นตัวเลข — ตัดช่องว่าง/ขึ้นบรรทัดใหม่ที่ pdfplumber
    บางทีแทรกกลางตัวเลขออกก่อน (เช่น '7 ,200' หรือมีบรรทัดอื่นปนมาด้วย '\\n-95,861')"""
    if cell is None:
        return None
    s = re.sub(r"\s+", "", str(cell))
    nums = re.findall(r"-?[\d,]+(?:\.\d+)?", s)
    return float(nums[-1].replace(",", "")) if nums else None


def parse_offo_pdf(url):
    """อ่านค่าจากตารางในรายงาน PDF โดยตรง (ไม่ใช้ regex บนข้อความดิบ ซึ่งพัง่ายเพราะ
    pdfplumber มักแทรกช่องว่าง/ตัดบรรทัดกลางตัวเลขตอนแปลงตารางเป็นข้อความ) ใช้คอลัมน์
    "รวม" (น้ำมัน+LPG รวมกัน ก่อนปรับด้วยเงินเรี่ยไร) เพราะเช็คกับตัวอย่างจริงแล้วตรงและ
    ไม่มีปัญหาการตัดบรรทัดแบบคอลัมน์ขวาสุด "ฐานะกองทุนรวม"."""
    r = requests.get(url, timeout=30)
    r.raise_for_status()
    if not url.lower().endswith(".pdf"):
        return None  # รายงานที่เป็นรูปภาพ (.png/.jpg) อ่านอัตโนมัติไม่ได้ ต้องกรอกมือ

    balance = cash = None
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                for row in table:
                    label = "".join(str(c) for c in row if c) if row else ""
                    if cash is None and "เงินฝากที่กรมบัญชีกลาง" in label:
                        cash = _num(row[6]) if len(row) > 6 else None
                    if balance is None and "ฐานะกองทุน" in label and "สุทธิ" in label:
                        balance = _num(row[6]) if len(row) > 6 else None

    return {"balance": balance, "cash": cash} if balance is not None and cash is not None else None


def runway_status(runway, settings):
    if runway is None:
        return "OK"
    if runway <= settings.get("critical_runway_days", 20):
        return "CRITICAL"
    if runway <= settings.get("watch_runway_days", 35):
        return "WATCH"
    return "OK"


def load_fund_settings(tok):
    rows = sheets_get(tok, SHEET_ID, f"{TABS['SETTING']}!J1:K10")
    out = {}
    key_map = {"critical runway days": "critical_runway_days", "watch runway days": "watch_runway_days"}
    for r in rows:
        if len(r) >= 2:
            k = key_map.get(str(r[0]).strip().lower())
            if k:
                try:
                    out[k] = float(r[1])
                except (TypeError, ValueError):
                    pass
    return out


def ingest_oilfund(tok, today: date, eppo_result):
    existing = col_to_dicts(sheets_get(tok, SHEET_ID, f"{TABS['ENTRY_OILFUND_SUSTAINABILITY']}!A1:H5000"))
    if not existing:
        print("[OILFUND] no existing rows to roll forward from, skipping")
        return
    last = existing[-1]
    last_date = excel_serial_to_date(last["Date"])
    if last_date is None:
        print("[OILFUND] could not parse last row's date, skipping")
        return

    if last_date == today:
        # แถวของวันนี้มีอยู่แล้ว — ถ้า Daily_Subsidy ยังว่าง (รันรอบก่อนหน้านี้ตอน EPPO ของวันนี้
        # ยังไม่เผยแพร่) แต่ตอนนี้มี eppo_result แล้ว ให้อัปเดตแถวเดิมให้ถูกต้องแทนที่จะปล่อยว่างค้างไว้
        subsidy_blank = last.get("Daily_Subsidy (ล้านบาท/วัน)") in (None, "")
        if subsidy_blank and eppo_result and len(existing) >= 2:
            prev = existing[-2]
            prev_balance = float(prev["Total_Balance (ล้านบาท)"])
            prev_cash = float(prev["Cash_Remaining (ล้านบาท)"])
            settings = load_fund_settings(tok)
            subsidy, collection, net_impact = eppo_result["subsidy"], eppo_result["collection"], eppo_result["net_impact"]
            balance = round(prev_balance + net_impact, 2)
            cash = round(prev_cash - subsidy, 2)
            runway = round(cash / subsidy, 2) if subsidy else None
            status = runway_status(runway, settings)
            row_num = len(existing) + 1  # +1 สำหรับแถวหัวตาราง
            sheets_update(tok, SHEET_ID, f"{TABS['ENTRY_OILFUND_SUSTAINABILITY']}!B{row_num}:H{row_num}",
                          [[balance, cash, round(subsidy, 2), round(collection, 2), round(net_impact, 2),
                            runway if runway is not None else "", status]])
            print(f"[OILFUND] {today}: filled in Daily_Subsidy/Collection/Runway now that EPPO data is available (was blank)")
        else:
            print(f"[OILFUND] already up to date ({last_date})")
        return
    if last_date > today:
        print(f"[OILFUND] already up to date ({last_date})")
        return

    settings = load_fund_settings(tok)
    report_end, report_url = latest_offo_report()
    official_date = report_end + timedelta(days=1) if report_end else None

    new_rows = []
    balance = float(last["Total_Balance (ล้านบาท)"])
    cash = float(last["Cash_Remaining (ล้านบาท)"])
    prev_subsidy = float(last["Daily_Subsidy (ล้านบาท/วัน)"])
    prev_net_impact = float(last["Net_Fund_Impact (ล้านบาท/วัน)"])

    # ถ้าห่างจากแถวล่าสุดเกิน 7 วัน (เช่น สคริปต์นี้ไม่เคยรันมานาน) อย่า backfill ทุกวันด้วยอัตรา
    # เก่าๆ ซ้ำๆ (ประมาณการเพี้ยนสะสม) — ข้ามช่วงที่ไม่มีข้อมูลไปเลย แล้วเริ่มนับใหม่จาก
    # รายงานทางการล่าสุด (ถ้ามีออกมาระหว่างนั้น) หรือไม่ก็วันนี้ โดยใช้ยอดล่าสุดที่มีเป็นฐาน
    gap_days = (today - last_date).days
    if gap_days > 7:
        if official_date and official_date > last_date:
            print(f"[OILFUND] gap of {gap_days} days — skipping ahead to the latest official report ({official_date}) instead of backfilling every day")
            d = official_date
        else:
            print(f"[OILFUND] gap of {gap_days} days and no newer official report — skipping ahead to today, carrying forward the last known balance as-is")
            d = today
    else:
        d = last_date + timedelta(days=1)
    while d <= today:
        date_str = d.isoformat()
        if official_date == d:
            parsed = parse_offo_pdf(report_url)
            if parsed:
                balance, cash = parsed["balance"], parsed["cash"]
                print(f"[OILFUND] {date_str}: using official report ({report_url})")
            else:
                print(f"[OILFUND] {date_str}: official report exists but couldn't be parsed (image or format change) — needs manual entry")
                d += timedelta(days=1)
                continue
            subsidy = collection = net_impact = None  # unknown for the reset day itself
        else:
            balance = balance + prev_net_impact
            cash = cash - prev_subsidy
            if d == today and eppo_result:
                subsidy, collection, net_impact = eppo_result["subsidy"], eppo_result["collection"], eppo_result["net_impact"]
            else:
                subsidy = collection = net_impact = None  # no EPPO data for this in-between day

        runway = round(cash / subsidy, 2) if subsidy else None
        status = runway_status(runway, settings)
        new_rows.append([
            date_str, round(balance, 2), round(cash, 2),
            round(subsidy, 2) if subsidy is not None else "",
            round(collection, 2) if collection is not None else "",
            round(net_impact, 2) if net_impact is not None else "",
            runway if runway is not None else "", status,
        ])
        if subsidy is not None:
            prev_subsidy, prev_net_impact = subsidy, net_impact
        d += timedelta(days=1)

    sheets_append(tok, SHEET_ID, f"{TABS['ENTRY_OILFUND_SUSTAINABILITY']}!A1:H1", new_rows)
    print(f"[OILFUND] appended {len(new_rows)} rows")


# ============================================================
# MAIN
# ============================================================
def main():
    tok = get_token()
    today = datetime.now(TZ).date()

    # แต่ละแหล่งข้อมูลเป็นอิสระจากกัน — ถ้าแหล่งหนึ่งล้ม (เว็บเปลี่ยนฟอร์แมต, ยังไม่เผยแพร่ ฯลฯ)
    # ไม่ควรทำให้แหล่งอื่นที่ยังทำงานได้ปกติถูกข้ามไปด้วย แต่ยังต้องรายงานว่าล้มเหลว
    # (ไม่ใช่แค่ print แล้วให้ script จบแบบ exit code 0 เหมือนไม่มีอะไรผิดพลาด)
    failures = []

    eppo_result = None
    try:
        eppo_result = ingest_eppo(tok, today)
    except Exception as e:
        print(f"[EPPO] FAILED: {e}")
        failures.append("EPPO")

    try:
        ingest_mops(tok, today)
    except Exception as e:
        print(f"[MOPS] FAILED: {e}")
        failures.append("MOPS")

    try:
        ingest_oilfund(tok, today, eppo_result)
    except Exception as e:
        print(f"[OILFUND] FAILED: {e}")
        failures.append("OILFUND")

    try:
        ingest_feeds(tok)
    except Exception as e:
        print(f"[FEED] FAILED: {e}")
        failures.append("FEED")

    if failures:
        raise SystemExit(f"ingest.py: {', '.join(failures)} failed — see logs above")


if __name__ == "__main__":
    main()
